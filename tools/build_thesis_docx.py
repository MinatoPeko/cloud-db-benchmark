#!/usr/bin/env python3
"""
tools/build_thesis_docx.py
Chuyển đổi toàn văn BAO_CAO_TOT_NGHIEP_PHENIKAA.md sang file Word:
- Tự động thay thế 6 khối mã Mermaid bằng 6 hình vẽ sơ đồ chuyên nghiệp (300 DPI).
- Tự động chèn 4 biểu đồ kết quả đo đạc từ results/charts/.
- XÓA SẠCH 100% các thẻ HTML như <br>, <br/>, <br><br> và chuyển thành ngắt dòng thực tế (\n).
- Dịch các ký hiệu toán học LaTeX (\le, \ge, \times...) thành ký tự Unicode chuẩn (≤, ≥, ×...).
- Căn chỉnh định dạng chuẩn Đại học Phenikaa: Times New Roman, lề A4, bảng biểu có viền & màu nền, căn đều hai bên (Justify).
"""

import os
import sys
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD_PATH = os.path.join(BASE_DIR, "BAO_CAO_TOT_NGHIEP_PHENIKAA.md")
DOCX_TARGETS = [
    os.path.join(BASE_DIR, "temp.docx"),
    os.path.join(BASE_DIR, "temp_v2.docx"),
    os.path.join(BASE_DIR, "temp_v3.docx")
]
DIAGRAMS_DIR = os.path.join(BASE_DIR, "results", "diagrams")
CHARTS_DIR = os.path.join(BASE_DIR, "results", "charts")

# Bảng màu chuẩn học thuật Phenikaa
COLOR_PRIMARY_NAVY = RGBColor(30, 58, 138)    # #1E3A8A - Tiêu đề Chương
COLOR_SECONDARY = RGBColor(15, 23, 42)       # #0F172A - Tiêu đề Mục
COLOR_BODY = RGBColor(30, 41, 59)            # #1E293B - Văn bản chính
COLOR_MUTED = RGBColor(100, 116, 139)        # #64748B - Chú thích
HEX_HEADER_BG = "EFF6FF"                     # Xanh nhạt cho tiêu đề bảng
HEX_BORDER = "CBD5E1"                        # Viền xám mảnh cho bảng


def set_cell_background(cell, hex_color):
    """Đặt màu nền cho ô bảng."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    """Đặt lề trong cho ô bảng (dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Đặt viền thanh mảnh cho toàn bộ bảng."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def clean_math_and_html(text):
    """Làm sạch ký tự HTML và chuyển đổi ký hiệu toán học LaTeX sang Unicode chuẩn."""
    # Chuyển đổi <br>, <br/>, <br /> thành dấu xuống dòng thực tế \n
    text = re.sub(r'<\s*br\s*/?>', '\n', text, flags=re.IGNORECASE)

    # Chuyển đổi LaTeX inline / display
    text = re.sub(r'\$\$(.*?)\$\$', r'\1', text)
    text = re.sub(r'\$(.*?)\$', r'\1', text)

    # Ký hiệu toán học LaTeX thông dụng
    math_replacements = [
        (r'\\le\b', '≤'),
        (r'\\ge\b', '≥'),
        (r'\\times\b', '×'),
        (r'\\approx\b', '≈'),
        (r'\\sim\b', '~'),
        (r'\\in\b', '∈'),
        (r'\\gamma\b', 'γ'),
        (r'\\Pi\b', 'Π'),
        (r'\\sum\b', '∑'),
        (r'\\iff\b', '⇔'),
        (r'\\text\{(.*?)\}', r'\1'),
        (r'\\log\b', 'log'),
    ]
    for pattern, rep in math_replacements:
        text = re.sub(pattern, rep, text)

    return text


def format_runs(p, text, font_size=None, code_font_size=None):
    """Phân tích các thẻ Markdown (**bold**, *italic*, `code`, \n) thành các run trong paragraph."""
    text = clean_math_and_html(text)

    # Chia nhỏ theo các dòng ngắt \n
    lines = text.split('\n')
    for line_idx, line in enumerate(lines):
        if line_idx > 0:
            p.add_run().add_break()  # Xuống dòng mềm trong Word (Shift + Enter)

        tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', line)
        for tok in tokens:
            if not tok:
                continue
            if tok.startswith('**') and tok.endswith('**'):
                r = p.add_run(tok[2:-2])
                r.bold = True
                if font_size:
                    r.font.size = font_size
            elif tok.startswith('*') and tok.endswith('*'):
                r = p.add_run(tok[1:-1])
                r.italic = True
                if font_size:
                    r.font.size = font_size
            elif tok.startswith('`') and tok.endswith('`'):
                r = p.add_run(tok[1:-1])
                r.font.name = 'Consolas'
                r.font.size = code_font_size if code_font_size else Pt(9.5)
                r.font.color.rgb = RGBColor(15, 23, 42)
            else:
                r = p.add_run(tok)
                if font_size:
                    r.font.size = font_size


def parse_markdown_table_row(line):
    """Tách các ô trong bảng Markdown, bảo vệ các ký tự | bên trong khối mã `code` hoặc escaped \|."""
    chars = []
    in_code = False
    escaped = False
    for c in line:
        if c == '\\' and not escaped:
            escaped = True
            continue
        if escaped:
            if c == '|':
                chars.append('__ESCAPED_PIPE__')
            else:
                chars.append('\\' + c)
            escaped = False
            continue
        if c == '`':
            in_code = not in_code
        if c == '|' and in_code:
            chars.append('__PIPE__')
        else:
            chars.append(c)
    masked = "".join(chars).strip()
    if masked.startswith("|"): masked = masked[1:]
    if masked.endswith("|"): masked = masked[:-1]
    cells = [
        c.replace('__PIPE__', '|').replace('__ESCAPED_PIPE__', '|').strip()
        for c in masked.split('|')
    ]
    return cells


def get_table_column_widths(header_row, num_cols):
    """Tính toán chiều rộng cột tối ưu cho từng bảng học thuật chuẩn khổ A4 (printable width ~ 6.30 in)."""
    hdr_text = " ".join(header_row)

    # Bảng 3.1: Ma trận truy vết Use Cases (6 cột)
    if "Mã UC" in hdr_text and "PostgreSQL" in hdr_text:
        return [Inches(0.55), Inches(1.05), Inches(1.55), Inches(1.50), Inches(1.10), Inches(0.55)]

    # Bảng 4.2: Bảng tổng hợp số liệu đo kiểm hiệu năng (10 cột)
    if "Mean (ms)" in hdr_text and "p99 (ms)" in hdr_text:
        return [Inches(0.80), Inches(0.45), Inches(0.90), Inches(0.48), Inches(0.44),
                Inches(0.48), Inches(0.48), Inches(0.55), Inches(0.48), Inches(0.60)]

    # Bảng thông tin sinh viên & GVHD ở bìa (2 cột)
    if "Thông tin sinh viên" in hdr_text:
        return [Inches(2.80), Inches(3.50)]

    # Danh mục từ viết tắt (3 cột)
    if "Chữ viết tắt" in hdr_text:
        return [Inches(1.20), Inches(2.10), Inches(3.00)]

    # Danh mục bảng biểu hoặc danh mục hình vẽ (3 cột)
    if "Ký hiệu bảng" in hdr_text or "Ký hiệu hình" in hdr_text:
        return [Inches(1.10), Inches(4.30), Inches(0.90)]

    # Bảng so sánh 3 hệ CSDL (4 cột: Tiêu chí/Nghiệp vụ, Postgres, Mongo, Neo4j)
    if num_cols == 4 and ("PostgreSQL" in hdr_text or "Tiêu chí" in hdr_text or "Mục đích" in hdr_text):
        return [Inches(1.50), Inches(1.60), Inches(1.60), Inches(1.60)]

    # Bảng 3.3: Tham số mô hình Barabási-Albert (7 cột)
    if "Quy mô" in hdr_text and num_cols == 7:
        return [Inches(0.80), Inches(0.90), Inches(0.80), Inches(1.00), Inches(1.00), Inches(1.00), Inches(0.80)]

    # Bảng 4.1: Đối soát toàn vẹn dữ liệu (6 cột)
    if "Thực thể dữ liệu" in hdr_text and num_cols == 6:
        return [Inches(1.40), Inches(1.00), Inches(1.00), Inches(1.00), Inches(0.90), Inches(1.00)]

    # Bảng 4.3: Tổng kết hệ CSDL thắng thế (5 cột)
    if "Mã Use Case" in hdr_text and "thắng thế" in hdr_text:
        return [Inches(0.70), Inches(1.35), Inches(1.35), Inches(1.35), Inches(1.55)]

    # Bảng 5.1: Ma trận khuyến nghị kiến trúc (5 cột)
    if ("Tiêu chí nghiệp vụ" in hdr_text or "Khuyến nghị" in hdr_text) and num_cols == 5:
        return [Inches(1.30), Inches(1.20), Inches(1.20), Inches(1.20), Inches(1.40)]

    # Mặc định: chia đều tổng chiều rộng khả dụng 6.30 inch
    w = Inches(6.30 / num_cols)
    return [w] * num_cols


def add_image_with_caption(doc, img_path, caption_text, width_inches=5.8):
    """Chèn ảnh căn giữa kèm dòng chú thích in nghiêng."""
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_inches))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(10)
        run_cap = p_cap.add_run(caption_text)
        run_cap.italic = True
        run_cap.font.size = Pt(10.5)
        run_cap.font.color.rgb = COLOR_MUTED
    else:
        p_err = doc.add_paragraph(f"[Không tìm thấy tệp ảnh: {os.path.basename(img_path)}]")
        p_err.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_docx():
    print(f"📖 Đang đọc tệp nguồn: {MD_PATH}")
    with open(MD_PATH, "r", encoding="utf-8") as f:
        content = f.read()

    doc = Document()

    # 1. Cấu hình lề trang chuẩn A4 theo quy định của Phenikaa
    for section in doc.sections:
        section.top_margin = Inches(0.79)     # 2.0 cm
        section.bottom_margin = Inches(0.79)  # 2.0 cm
        section.left_margin = Inches(1.18)    # 3.0 cm (đóng gáy sách)
        section.right_margin = Inches(0.79)   # 2.0 cm

    # 2. Cấu hình Style mặc định
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(12)
    normal_font.color.rgb = COLOR_BODY

    lines = content.split('\n')
    i = 0
    total_lines = len(lines)
    mermaid_count = 0

    in_mermaid = False
    mermaid_block_lines = []

    in_table = False
    table_raw_rows = []

    in_code_block = False
    code_block_lines = []
    seen_first_page_break = False

    while i < total_lines:
        line = lines[i]
        stripped = line.strip()

        # -------------------------------------------------------------
        # XỬ LÝ CÁC DÒNG CHỈ CÓ THẺ <br>
        # -------------------------------------------------------------
        if re.match(r'^(<\s*br\s*/?>\s*)+$', stripped, flags=re.IGNORECASE):
            # Bỏ qua hoàn toàn dòng chỉ chứa thẻ <br>, không tạo đoạn văn thừa
            i += 1
            continue

        # -------------------------------------------------------------
        # XỬ LÝ MERMAID BLOCKS
        # -------------------------------------------------------------
        if stripped.startswith("```mermaid"):
            in_mermaid = True
            mermaid_block_lines = []
            i += 1
            continue

        if in_mermaid:
            if stripped.startswith("```"):
                in_mermaid = False
                mermaid_count += 1
                mermaid_text = "\n".join(mermaid_block_lines)

                if "1. Nghiên cứu Lý thuyết" in mermaid_text or mermaid_count == 1:
                    img_path = os.path.join(DIAGRAMS_DIR, "hinh_1_1_phuong_phap_luan.png")
                    cap = "Hình 1.1: Sơ đồ luồng phương pháp luận nghiên cứu thực nghiệm 5 bước"
                    add_image_with_caption(doc, img_path, cap, width_inches=5.2)
                elif "Paradigms" in mermaid_text or mermaid_count == 2:
                    img_path = os.path.join(DIAGRAMS_DIR, "hinh_2_1_3_paradigms.png")
                    cap = "Hình 2.1: So sánh đặc tính 3 trường phái kiến trúc Cơ sở Dữ liệu"
                    add_image_with_caption(doc, img_path, cap, width_inches=6.0)
                elif "EngineComparison" in mermaid_text or mermaid_count == 3:
                    img_path = os.path.join(DIAGRAMS_DIR, "hinh_2_2_engine_comparison.png")
                    cap = "Hình 2.2: So sánh kiến trúc Động cơ Lưu trữ & Truy xuất (PostgreSQL, MongoDB, Neo4j)"
                    add_image_with_caption(doc, img_path, cap, width_inches=6.0)
                elif "DevContainerStack" in mermaid_text or mermaid_count == 4:
                    img_path = os.path.join(DIAGRAMS_DIR, "hinh_2_3_devcontainer_stack.png")
                    cap = "Hình 2.3: Kiến trúc hạ tầng ảo hóa Cloud Dev Container trên GitHub Codespaces"
                    add_image_with_caption(doc, img_path, cap, width_inches=5.8)
                elif "erDiagram" in mermaid_text or mermaid_count == 5:
                    img_path = os.path.join(DIAGRAMS_DIR, "hinh_3_1_erd_social_network.png")
                    cap = "Hình 3.1: Sơ đồ thực thể liên kết (ERD) miền bài toán Mạng xã hội"
                    add_image_with_caption(doc, img_path, cap, width_inches=6.0)
                elif "sequenceDiagram" in mermaid_text or mermaid_count == 6:
                    img_path = os.path.join(DIAGRAMS_DIR, "hinh_3_5_sequence_benchmark.png")
                    cap = "Hình 3.5: Trình tự thực thi đo kiểm và giám sát tài nguyên của Benchmark Harness"
                    add_image_with_caption(doc, img_path, cap, width_inches=6.0)
                
                i += 1
                continue
            else:
                mermaid_block_lines.append(line)
                i += 1
                continue

        # -------------------------------------------------------------
        # XỬ LÝ KHỐI CODE BÌNH THƯỜNG
        # -------------------------------------------------------------
        if stripped.startswith("```") and not in_code_block:
            in_code_block = True
            code_block_lines = []
            i += 1
            continue

        if in_code_block:
            if stripped.startswith("```"):
                in_code_block = False
                code_text = "\n".join(code_block_lines)
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                cell.width = Inches(6.3)
                set_cell_background(cell, "F8FAFC")
                set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
                p_code = cell.paragraphs[0]
                p_code.paragraph_format.space_before = Pt(0)
                p_code.paragraph_format.space_after = Pt(0)
                p_code.paragraph_format.line_spacing = 1.15
                r_code = p_code.add_run(code_text)
                r_code.font.name = 'Consolas'
                r_code.font.size = Pt(9.5)
                r_code.font.color.rgb = RGBColor(15, 23, 42)
                
                p_space = doc.add_paragraph()
                p_space.paragraph_format.space_before = Pt(0)
                p_space.paragraph_format.space_after = Pt(6)
                
                i += 1
                continue
            else:
                code_block_lines.append(line)
                i += 1
                continue

        # -------------------------------------------------------------
        # XỬ LÝ BẢNG BIỂU MARKDOWN
        # -------------------------------------------------------------
        if stripped.startswith("|") and stripped.endswith("|"):
            if not in_table:
                in_table = True
                table_raw_rows = []
            table_raw_rows.append(stripped)
            i += 1
            continue

        if in_table and not (stripped.startswith("|") and stripped.endswith("|")):
            in_table = False
            cleaned_rows = []
            for r in table_raw_rows:
                cells = parse_markdown_table_row(r)
                if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                    continue
                cleaned_rows.append(cells)

            if cleaned_rows:
                num_rows = len(cleaned_rows)
                num_cols = max(len(r) for r in cleaned_rows)
                
                t = doc.add_table(rows=num_rows, cols=num_cols)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(t, color=HEX_BORDER, sz="4")

                col_widths = get_table_column_widths(cleaned_rows[0], num_cols)
                hdr_str = " ".join(cleaned_rows[0])
                is_traceability_table = ("Mã UC" in hdr_str and "PostgreSQL" in hdr_str)
                is_metric_table = ("Mean (ms)" in hdr_str and "p99 (ms)" in hdr_str)

                for row_idx, row_data in enumerate(cleaned_rows):
                    row_obj = t.rows[row_idx]
                    is_header = (row_idx == 0)

                    # Ngăn chặn ngắt dòng giữa trang và lặp lại tiêu đề bảng
                    trPr = row_obj._tr.get_or_add_trPr()
                    trPr.append(OxmlElement('w:cantSplit'))
                    if is_header:
                        trPr.append(OxmlElement('w:tblHeader'))

                    for col_idx in range(num_cols):
                        cell_obj = row_obj.cells[col_idx]
                        if col_idx < len(col_widths):
                            cell_obj.width = col_widths[col_idx]

                        # Cấu hình padding ô phù hợp với tính chất bảng
                        if is_traceability_table:
                            set_cell_margins(cell_obj, top=35, bottom=35, left=50, right=50)
                        elif is_metric_table:
                            set_cell_margins(cell_obj, top=30, bottom=30, left=35, right=35)
                        else:
                            set_cell_margins(cell_obj, top=50, bottom=50, left=70, right=70)
                        
                        val = row_data[col_idx] if col_idx < len(row_data) else ""
                        p_cell = cell_obj.paragraphs[0]
                        p_cell.paragraph_format.space_before = Pt(1)
                        p_cell.paragraph_format.space_after = Pt(1)
                        p_cell.paragraph_format.line_spacing = 1.05

                        if is_header:
                            set_cell_background(cell_obj, HEX_HEADER_BG)
                            p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            format_runs(p_cell, val, font_size=Pt(8.5 if is_metric_table else 9.5))
                            for run in p_cell.runs:
                                run.bold = True
                                run.font.color.rgb = COLOR_PRIMARY_NAVY
                        else:
                            if is_traceability_table:
                                if col_idx in [0, 5]:
                                    p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                else:
                                    p_cell.alignment = WD_ALIGN_PARAGRAPH.LEFT
                                format_runs(p_cell, val, font_size=Pt(8.5), code_font_size=Pt(8.0))
                            elif is_metric_table:
                                if col_idx in [0, 1, 2]:
                                    p_cell.alignment = WD_ALIGN_PARAGRAPH.LEFT
                                else:
                                    p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                format_runs(p_cell, val, font_size=Pt(8.0))
                            else:
                                p_cell.alignment = WD_ALIGN_PARAGRAPH.LEFT
                                format_runs(p_cell, val, font_size=Pt(9.5), code_font_size=Pt(8.5))

                p_after_tbl = doc.add_paragraph()
                p_after_tbl.paragraph_format.space_after = Pt(6)

        # -------------------------------------------------------------
        # XỬ LÝ NGẮT TRANG (\newpage)
        # -------------------------------------------------------------
        if stripped == "\\newpage":
            doc.add_page_break()
            seen_first_page_break = True
            i += 1
            continue

        # -------------------------------------------------------------
        # XỬ LÝ DÒNG TIÊU ĐỀ (#, ##, ###)
        # -------------------------------------------------------------
        if stripped.startswith("# "):
            heading_text = stripped[2:].strip()
            if "BÁO CÁO ĐỒ ÁN TỐT NGHIỆP" in heading_text or "CHƯƠNG" in heading_text:
                if i > 5 and seen_first_page_break:
                    doc.add_page_break()

            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True

            if not seen_first_page_break:
                # Trang bìa
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(heading_text)
                r.bold = True
                r.font.size = Pt(16 if "BÁO CÁO" in heading_text else 14)
                r.font.color.rgb = COLOR_PRIMARY_NAVY
            else:
                if "CHƯƠNG" in heading_text or "TÀI LIỆU" in heading_text or "PHỤ LỤC" in heading_text:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(heading_text)
                r.bold = True
                r.font.size = Pt(15)
                r.font.color.rgb = COLOR_PRIMARY_NAVY
            i += 1
            continue

        if stripped.startswith("## "):
            heading_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True

            if not seen_first_page_break:
                # Trang bìa
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(heading_text)
                r.bold = True
                r.font.size = Pt(14 if "NGHIÊN CỨU" in heading_text else 13)
                r.font.color.rgb = COLOR_PRIMARY_NAVY
            else:
                if any(k in heading_text for k in ["LỜI CAM ĐOAN", "LỜI CẢM ƠN", "DANH MỤC", "TÀI LIỆU THAM KHẢO", "PHỤ LỤC"]):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(14)
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(heading_text)
                r.bold = True
                r.font.size = Pt(13)
                r.font.color.rgb = COLOR_PRIMARY_NAVY
            i += 1
            continue

        if stripped.startswith("### "):
            heading_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True

            if not seen_first_page_break:
                # Trang bìa
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(heading_text)
                r.bold = True
                r.font.size = Pt(12)
                r.font.color.rgb = COLOR_SECONDARY
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(heading_text)
                r.bold = True
                r.font.size = Pt(12)
                r.font.color.rgb = COLOR_SECONDARY
            i += 1
            continue

        # -------------------------------------------------------------
        # XỬ LÝ CHÈN 4 BIỂU ĐỒ KẾT QUẢ ĐO ĐẠC Ở CHƯƠNG 4
        # -------------------------------------------------------------
        if "01_latency_bar_small_" in stripped:
            chart_file = os.path.join(CHARTS_DIR, "01_latency_bar_small_20261024_233530.png")
            add_image_with_caption(doc, chart_file, "Hình 4.1: Biểu đồ so sánh Thời gian đáp ứng trung bình (Mean Latency) giữa 3 CSDL", width_inches=5.8)
            i += 1
            continue

        if "02_p99_heatmap_small_" in stripped:
            chart_file = os.path.join(CHARTS_DIR, "02_p99_heatmap_small_20261024_233530.png")
            add_image_with_caption(doc, chart_file, "Hình 4.2: Bản đồ nhiệt (Heatmap) thể hiện độ trễ phân vị rủi ro đuôi p99 (Worst-case Latency)", width_inches=5.8)
            i += 1
            continue

        if "03_qps_bar_small_" in stripped:
            chart_file = os.path.join(CHARTS_DIR, "03_qps_bar_small_20261024_233530.png")
            add_image_with_caption(doc, chart_file, "Hình 4.3: Biểu đồ so sánh Thông lượng xử lý truy vấn trên giây (Throughput QPS)", width_inches=5.8)
            i += 1
            continue

        if "04_ram_bar_small_" in stripped:
            chart_file = os.path.join(CHARTS_DIR, "04_ram_bar_small_20261024_233530.png")
            add_image_with_caption(doc, chart_file, "Hình 4.4: Biểu đồ so sánh Mức tiêu thụ bộ nhớ RAM Đỉnh (RAM Peak Footprint MB)", width_inches=5.8)
            i += 1
            continue

        if stripped == "---" or stripped == "***":
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        # Dòng chú thích bảng biểu: *Bảng X.Y: ...*
        if re.match(r'^\*Bảng\s+\d+\.\d+:', stripped):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            format_runs(p, stripped, font_size=Pt(10.5))
            for run in p.runs:
                run.bold = True
                run.italic = True
                run.font.color.rgb = COLOR_SECONDARY
            i += 1
            continue

        # Khối chữ ký và ngày tháng
        if any(sig in stripped for sig in ["Hà Nội, ngày 24", "Sinh viên thực hiện", "(Ký và ghi rõ họ tên)", "Vũ Đức Hiếu"]):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.right_indent = Inches(0.5)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            format_runs(p, stripped)
            i += 1
            continue

        # Trang bìa: text căn giữa (như Hà Nội – Năm 2026)
        if not seen_first_page_break:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            format_runs(p, stripped)
            i += 1
            continue

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25

        if stripped.startswith("- ") or stripped.startswith("+ ") or stripped.startswith("* "):
            p.paragraph_format.left_indent = Inches(0.25)
            format_runs(p, stripped[2:].strip())
        elif re.match(r'^\d+\.\s+', stripped):
            p.paragraph_format.left_indent = Inches(0.25)
            format_runs(p, stripped)
        else:
            format_runs(p, stripped)

        i += 1

    # -------------------------------------------------------------
    # LƯU FILE VÀ TẠO FILE CÓ KHUNG VIỀN ĐÔI (DOUBLE PAGE BORDER)
    # -------------------------------------------------------------
    # 1. Tệp chính duy nhất của toàn bộ đồ án: BAO_CAO_TOT_NGHIEP_PHENIKAA.docx
    thesis_path = os.path.join(BASE_DIR, "BAO_CAO_TOT_NGHIEP_PHENIKAA.docx")
    try:
        doc.save(thesis_path)
        file_size_kb = os.path.getsize(thesis_path) / 1024
        print(f"🎉 ĐÃ LƯU BÁO CÁO CHÍNH: {os.path.basename(thesis_path)} ({file_size_kb:.1f} KB)")

        # Thêm khung viền đôi trực tiếp vào tệp báo cáo chính
        sys.path.append(BASE_DIR)
        from edit_page_border import add_double_border
        add_double_border(thesis_path)
        print(f"🎉 ĐÃ ÁP DỤNG KHUNG VIỀN ĐÔI VÀO: {os.path.basename(thesis_path)} ({os.path.getsize(thesis_path)/1024:.1f} KB)")
    except PermissionError:
        print(f"⚠️ {os.path.basename(thesis_path)} đang mở trong Word, vui lòng đóng để cập nhật!")

    # 2. Đồng bộ sang các file temp để dự phòng
    temp_files = [
        os.path.join(BASE_DIR, "temp_v3.docx"),
        os.path.join(BASE_DIR, "temp_v3_fixed.docx"),
    ]
    for tf in temp_files:
        try:
            doc.save(tf)
        except Exception:
            pass

    # 3. Tạo temp_v3_bordered.docx nếu có thể
    bordered_v3 = os.path.join(BASE_DIR, "temp_v3_bordered.docx")
    try:
        from edit_page_border import add_double_border
        add_double_border(thesis_path, bordered_v3)
    except Exception:
        pass


if __name__ == "__main__":
    build_docx()
